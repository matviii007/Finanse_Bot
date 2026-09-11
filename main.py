import asyncio
from aiogram import Bot, Dispatcher, F, Router, types
from aiogram.filters import Command, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
    BufferedInputFile
)
from aiogram.utils.keyboard import InlineKeyboardBuilder
from dotenv import load_dotenv
from datetime import datetime
import os
import asyncpg
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

evn_path = os.path.join(os.path.dirname(__file__), ".env")
load_dotenv(dotenv_path=evn_path)

DATABASE_URL = os.getenv("DATABASE_URL")
BOT_TOKEN = os.getenv("BOT_TOKEN")

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()

# --- КЛАВІАТУРИ ---
inline_kb_builder = InlineKeyboardBuilder()
inline_kb_builder.row(InlineKeyboardButton(text="Українська мова 🇺🇦", callback_data="lang_ua"))
inline_kb_builder.row(InlineKeyboardButton(text="English language 🇬🇧", callback_data="lang_en"))

inline_kb_ua = InlineKeyboardBuilder()
inline_kb_ua.row(InlineKeyboardButton(text="Переглянути табличку витрат та доходів", callback_data="view_table_ua"))
inline_kb_ua.row(InlineKeyboardButton(text="Оберіть мову", callback_data="saints_ua"))

inline_kb_en = InlineKeyboardBuilder()
inline_kb_en.row(InlineKeyboardButton(text="Check the income and expense table", callback_data="view_table_en"))
inline_kb_en.row(InlineKeyboardButton(text="Choose a language", callback_data="saints_en"))

menu_uk = ReplyKeyboardMarkup(
    keyboard=[
        [
            KeyboardButton(text="➕Додати витрати 💸"),
            KeyboardButton(text="➕Додати дохід 💵"),
        ],
        [
            KeyboardButton(text="Налаштування ⚙️"),
        ],
    ],
    resize_keyboard=True,
)

menu_en = ReplyKeyboardMarkup(
    keyboard=[
        [
            KeyboardButton(text="➕Add Expense 💸"),
            KeyboardButton(text="➕Add Income 💵"),
        ],
        [
            KeyboardButton(text="Settings ⚙️"),
        ],
    ],
    resize_keyboard=True,
)

comfirm_kb_ua = InlineKeyboardMarkup(
    inline_keyboard=[
        [
            InlineKeyboardButton(text="✅Підтвердити", callback_data="confirm"),
            InlineKeyboardButton(text="❌Скасувати", callback_data="cancel"),
        ]
    ]
)

comfirm_kb_en = InlineKeyboardMarkup(
    inline_keyboard=[
        [
            InlineKeyboardButton(text="✅Confirm", callback_data="confirm"),
            InlineKeyboardButton(text="❌Cancel", callback_data="cancel"),
        ]
    ]
)

file_action_kb_ua = InlineKeyboardMarkup(
    inline_keyboard=[
        [
            InlineKeyboardButton(text="Очистити історію 🗑️", callback_data="request_clear_all")
        ]
    ]
)

clear_confirm_kb_ua = InlineKeyboardMarkup(
    inline_keyboard=[
        [
            InlineKeyboardButton(text="✅ Так, очистити все", callback_data="confirm_clear_all"),
            InlineKeyboardButton(text="❌ Скасувати", callback_data="cancel_clear"),
        ]
    ]
)

class FinanceForm(StatesGroup):
    waiting_for_expense = State()
    waiting_for_expense_comment = State()
    waiting_for_income = State()
    waiting_for_income_comment = State()
    in_settings = State()
    waiting_for_confirmation = State()

# --- СТВОРЕННЯ EXCEL ---
async def generate_excel_and_stats(user_id: int, db_pool):
    async with db_pool.acquire() as conn:
        rows = await conn.fetch("""
            SELECT amount, is_income, comment, created_at 
            FROM operations 
            WHERE user_id = $1 AND created_at >= CURRENT_DATE - INTERVAL '180 days'
            ORDER BY created_at ASC;
        """, user_id)
    
    if not rows:
        return None, None

    max_expense = 0.0
    max_income = 0.0
    for row in rows:
        amt = float(row['amount'])
        if row['is_income']:
            if amt > max_income:
                max_income = amt
        else:
            if amt > max_expense:
                max_expense = amt

    wb = Workbook()
    ws = wb.active
    ws.title = "Фінанси"

    headers = ["Дата ", "Коментар", "витрати", "доходи", "В + чи -"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        created_dt = row['created_at']
        if isinstance(created_dt, datetime):
            date_str = created_dt.strftime("%d.%m.%Y %H:%M")
        else:
            date_str = str(created_dt)

        comment = row['comment'] or ""
        amount = float(row['amount'])
        
        if row['is_income']:
            expense_val = 0
            income_val = amount
            emoji_balance = f"🟢 🔼 +{amount:.2f}"
        else:
            expense_val = amount
            income_val = 0
            emoji_balance = f"🔴 🔽 -{amount:.2f}"
            
        ws.append([date_str, comment, expense_val, income_val, emoji_balance])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    stats_text = (
        "📊 Твоя фінансова статистика за пів року:\n\n"
        f"🟢 Найбільший дохід: {max_income:.2f} грн\n"
        f"🔴 Найбільша витрата: {max_expense:.2f} грн\n\n"
        "📁 Повний звіт у прикріпленому Excel-файлі нижче:"
    )

    excel_file = BufferedInputFile(file_stream.read(), filename="finance_report_6m.xlsx")
    return excel_file, stats_text

# --- ХЕНДЛЕРИ ---
@dp.message(F.text == "Переглянути табличку витрат та доходів")
@dp.message(F.text == "/table")
async def cmd_table(message: types.Message, db_pool):
    try:
        excel_file, stats_text = await generate_excel_and_stats(message.from_user.id, db_pool)
        if not excel_file:
            await message.answer("📊 У тебе поки немає збережених записів за останні 6 місяців.")
            return

        await message.answer_document(
            document=excel_file,
            caption=stats_text,
            parse_mode="Markdown",
            reply_markup=file_action_kb_ua
        )
    except Exception as e:
        print(f"❌ Помилка: {e}")
        await message.answer("Не вдалося сформувати Excel-табличку.")

@dp.callback_query(F.data == "view_table_ua")
async def process_callback_table_ua(callback: CallbackQuery, db_pool):
    await callback.answer()
    try:
        excel_file, stats_text = await generate_excel_and_stats(callback.from_user.id, db_pool)
        if not excel_file:
            await callback.message.answer("📊 У тебе поки немає збережених записів за останні 6 місяців.")
            return
        
        await callback.message.answer_document(
            document=excel_file, 
            caption=stats_text, 
            parse_mode="Markdown",
            reply_markup=file_action_kb_ua
        )
    except Exception as e:
        print(f"❌ Помилка: {e}")
        await callback.message.answer("Не вдалося сформувати Excel-табличку.")

@dp.callback_query(F.data == "request_clear_all")
async def process_request_clear(callback: CallbackQuery):
    await callback.answer()
    await callback.message.answer(
        "⚠️ УВАГА! Ви впевнені, що хочете повністю очистити історію та почати нову таблицю?\n"
        "Усі поточні записи буде видалено!",
        reply_markup=clear_confirm_kb_ua,
        parse_mode="Markdown"
    )
@dp.callback_query(F.data == "confirm_clear_all")
async def process_confirm_clear(callback: CallbackQuery, db_pool):
    await callback.answer()
    user_id = callback.from_user.id
    try:
        async with db_pool.acquire() as conn:
            await conn.execute("DELETE FROM operations WHERE user_id = $1;", user_id)
        await callback.message.edit_text("✨ Історію успішно очищено! Наступна табличка формуватиметься з чистого аркуша.")
    except Exception as e:
        print(f"❌ Помилка очищення: {e}")
        await callback.message.edit_text("Не вдалося очистити історію.")

@dp.callback_query(F.data == "cancel_clear")
async def process_cancel_clear(callback: CallbackQuery):
    await callback.answer()
    await callback.message.edit_text("Очищення скасовано. Ваші дані збережені 👍")

@dp.callback_query(F.data == "lang_ua")
async def set_ukrainian(callback: CallbackQuery):
    await callback.message.edit_text("Вибрано Українську мову 🇺🇦")
    await callback.message.answer('Привіт ! Я Бот для обліку фінансів 💰', reply_markup=menu_uk)

@dp.callback_query(F.data == "lang_en")
async def set_english(callback: CallbackQuery):
    await callback.message.edit_text("English language selected 🇬🇧")
    await callback.message.answer('Hello! I\'m your finances tracking bot 💰', reply_markup=menu_en)

@dp.message(CommandStart())
async def start_cmd(message: Message, state: FSMContext):
    await state.clear()
    remove_keyboard = await message.answer("...", reply_markup=ReplyKeyboardRemove())
    await remove_keyboard.delete()
    await message.answer("🌐 Оберіть мову / 🌐 Choose a language :", reply_markup=inline_kb_builder.as_markup())

@dp.message(Command("settings"))
@dp.message(F.text.contains("Налаштування ⚙️") | F.text.contains("Settings ⚙️"))
async def process_settings_button(message: Message, state: FSMContext):
    await state.set_state(FinanceForm.in_settings)
    await message.answer("⚙️ Ви обрали налаштування!", reply_markup=inline_kb_ua.as_markup())

@dp.message(F.text.contains("Додати витрати") | F.text.contains("Add Expense"))
async def process_expense(message: Message, state: FSMContext):
    await state.set_state(FinanceForm.waiting_for_expense)
    await message.answer("Введіть суму витрат 💸")

@dp.message(F.text.contains("Додати дохід") | F.text.contains("Add Income"))
async def process_income(message: Message, state: FSMContext):
    await state.set_state(FinanceForm.waiting_for_income)
    await message.answer("Введіть суму доходу 💵")

@dp.message(FinanceForm.waiting_for_expense)
async def filter_expense(message: Message, state: FSMContext):
    try:
        expense_amount = float(message.text.replace(',', '.'))
        if expense_amount <= 0:
            await message.answer("Введіть коректну суму витрат більше 0 💸")
            return
        await state.update_data(expense_amount=expense_amount)
        await state.set_state(FinanceForm.waiting_for_expense_comment)
        await message.answer("Введіть коментар до витрати (або відправте '-', якщо без коментаря) 📝:")
    except (TypeError, ValueError):
        await message.answer("Введіть коректну суму витрат 💸")

@dp.message(FinanceForm.waiting_for_expense_comment)
async def filter_expense_comment(message: Message, state: FSMContext):
    comment = "" if message.text.strip() == "-" else message.text.strip()
    await state.update_data(expense_comment=comment)
    data = await state.get_data()
    expense_amount = data.get("expense_amount")
    await state.set_state(FinanceForm.waiting_for_confirmation)
    await message.answer(
        f"Ви ввели суму витрат: {expense_amount} грн.\nКоментар: '{comment}'.\nБажаєте підтвердити чи скасувати?",
        reply_markup=comfirm_kb_ua
    )

@dp.message(FinanceForm.waiting_for_income)
async def filter_income(message: Message, state: FSMContext):
    try:
        income_amount = float(message.text.replace(',', '.'))
        if income_amount <= 0:
            await message.answer("Введіть коректну суму доходу більше 0 💵")
            return
        await state.update_data(income_amount=income_amount)
        await state.set_state(FinanceForm.waiting_for_income_comment)
        await message.answer("Введіть коментар до доходу (або відправте '-', якщо без коментаря) 📝:")
    except (TypeError, ValueError):
        await message.answer("Введіть коректну суму доходу 💵")

@dp.message(FinanceForm.waiting_for_income_comment)
async def filter_income_comment(message: Message, state: FSMContext):
    comment = "" if message.text.strip() == "-" else message.text.strip()
    await state.update_data(income_comment=comment)
    data = await state.get_data()
    income_amount = data.get("income_amount")
    await state.set_state(FinanceForm.waiting_for_confirmation)
    await message.answer(
        f"Ви ввели суму доходу: {income_amount} грн.\nКоментар: '{comment}'.\nБажаєте підтвердити чи скасувати?",
        reply_markup=comfirm_kb_ua
    )

@dp.callback_query(FinanceForm.waiting_for_confirmation, F.data == "confirm")
async def process_confirm_amount(callback: CallbackQuery, state: FSMContext, db_pool):
    await callback.answer()
    data = await state.get_data()
    user_id = callback.from_user.id
    
    if "expense_amount" in data:
        amount = data.get("expense_amount")
        comment = data.get("expense_comment", "")
        async with db_pool.acquire() as conn:
            await conn.execute("""
                INSERT INTO operations (user_id, amount, is_income, comment, created_at)
                VALUES ($1, $2, $3, $4, NOW());
            """, user_id, amount, False, comment)
        await callback.message.edit_text(f"Дані збережені! Витрата: {amount} грн ✅")
        
    elif "income_amount" in data:
        amount = data.get("income_amount")
        comment = data.get("income_comment", "")
        async with db_pool.acquire() as conn:
            await conn.execute("""
                INSERT INTO operations (user_id, amount, is_income, comment, created_at)
                VALUES ($1, $2, $3, $4, NOW());
            """, user_id, amount, True, comment)
        await callback.message.edit_text(f"Дані збережені! Дохід: {amount} грн ✅")
        
    await state.clear()

@dp.callback_query(FinanceForm.waiting_for_confirmation, F.data == "cancel")
async def process_cancel_amount(callback: CallbackQuery, state: FSMContext):
    await callback.answer()
    await state.clear()
    await callback.message.edit_text("Дію скасовано ⚙️")

async def get_db_pool():
    return await asyncpg.create_pool(
        dsn=DATABASE_URL,
        min_size=1,
        max_size=10
    )

async def main():
    db_pool = await get_db_pool()
    dp.workflow_data["db_pool"] = db_pool
    await dp.start_polling(bot)

if __name__ == "__main__":
    try:
        print("Бот запускається...")
        asyncio.run(main()) # Переконайся, що тут є дужки main()
    except Exception as e:
        print(f"❌ ВИНИКЛА ПОМИЛКА: {e}")
        input("\nНатисни Enter, щоб закрити...")